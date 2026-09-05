from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Thesis_Abstract.docx"


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ABSTRACT")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(31, 78, 121)
    title.paragraph_format.space_after = Pt(14)

    abstract = (
        "Accurate identification of promoter sequences is essential for genome annotation and the analysis of transcriptional regulation, yet it remains difficult because promoter architecture varies substantially across organisms and biological domains. This study developed a reproducible, leakage-aware computational framework for distinguishing promoter from non-promoter DNA sequences across Bacteria, Archaea, and Eukaryota. The finalized dataset comprised 99,321 sequences from Escherichia coli, Bacillus subtilis, selected archaeal organisms, Homo sapiens, Mus musculus, and Arabidopsis thaliana, including 49,094 promoters and 50,227 non-promoters. Sequence redundancy was reduced using CD-HIT-EST at 90% identity, and the finalized FASTA files were comprehensively audited. To prevent information leakage, exact sequences and reverse complements were assigned common group identifiers before deterministic 70:15:15 training, validation, and test partitioning. Six conflicting archaeal groups, comprising twelve records, were quarantined. "
        "DNA sequences were represented using one-hot nucleotide encoding for neural learning and complementary k-mer, position-specific, and dinucleotide-stability descriptors for feature-based classification. Multi-scale convolutional neural networks were finalized for the bacterial and archaeal domains, whereas a Residual CNN was finalized for Eukaryota. Each neural model was trained with three random seeds, and its probabilities were averaged. Organism-specific Random Forest models were then integrated with the neural ensembles using validation-selected probability weights and decision thresholds. Performance was evaluated on untouched test data using Matthews correlation coefficient (MCC), sensitivity, specificity, balanced accuracy, ROC-AUC, PR-AUC, Brier score, confusion matrices, and bootstrap confidence intervals. "
        "The three-seed neural ensembles obtained test MCC values of 0.617, 0.842, and 0.708 for Bacteria, Archaea, and Eukaryota, respectively. Hybridization increased the corresponding MCC values to 0.624, 0.844, and 0.723. Final hybrid ROC-AUC values were 0.879, 0.977, and 0.918, while PR-AUC values were 0.869, 0.965, and 0.939. Archaea achieved the strongest absolute performance, whereas Eukaryota obtained the largest MCC improvement from hybridization. Feature-ablation analysis further demonstrated that the most useful engineered representations differed among organisms. "
        "The findings show that promoter prediction benefits from combining domain specialization, multi-seed neural averaging, and taxon-specific engineered features. Equally, they demonstrate the importance of redundancy control, reverse-complement-aware partitioning, validation-only selection, and untouched-test evaluation. The resulting framework provides a reproducible foundation for multi-domain promoter prioritization, although broader species coverage, independent external validation, standardized negative controls, and experimental confirmation are required before universal application."
    )

    p = doc.add_paragraph(abstract)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(12)

    key = doc.add_paragraph()
    key.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    key.paragraph_format.line_spacing = 1.15
    key.add_run("Keywords: ").bold = True
    key.add_run(
        "promoter prediction; transcription start site; convolutional neural network; "
        "Random Forest; hybrid learning; multi-species genomics; sequence classification; "
        "information leakage; reverse complement; ensemble learning"
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Abstract | Domain-Aware Promoter Prediction")

    doc.core_properties.title = "Thesis Abstract"
    doc.core_properties.subject = "Domain-aware promoter prediction"
    doc.save(OUTPUT)
    print(OUTPUT)
    print("Abstract words:", len(re.findall(r"\b[\w–-]+\b", abstract)))


if __name__ == "__main__":
    build()
