from __future__ import annotations

import csv
import json
import os
from collections import Counter
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", "/tmp/matplotlib-promoter-report")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
FIGURES = ROOT / "results" / "main_figures"
RESULTS = ROOT / "results" / "main_tables" / "multiseed_ensemble.json"
MANIFEST = ROOT / "data" / "splits" / "split_manifest.csv"
OUTPUT = REPORTS / "Promoter_Prediction_Modeling_Flow_and_Final_Results.docx"

DOMAIN_LABELS = {"bacteria": "Bacteria", "archaea": "Archaea", "eukaryota": "Eukaryota"}
DOMAIN_COLORS = {"bacteria": "#247BA0", "archaea": "#D1495B", "eukaryota": "#2A9D6F"}


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    tc_pr.append(node)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text))
    run.bold = bold
    run.font.size = Pt(8.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF")
        shade(table.rows[0].cells[i], "24435C")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if len(table.rows) % 2 == 1:
                shade(cells[i], "EAF0F4")
    table.autofit = True
    return table


def load_data():
    result = json.loads(RESULTS.read_text())
    counts = Counter()
    with MANIFEST.open() as handle:
        for row in csv.DictReader(handle):
            counts[(row["domain"], row["class"], row["partition"])] += 1
    return result, counts


def workflow_figure(path):
    fig, ax = plt.subplots(figsize=(10.5, 3.3))
    ax.axis("off")
    labels = [
        "Finalized FASTA\nread-only audit",
        "Conflict quarantine\nand grouped 70/15/15 split",
        "Masked one-hot\nencoding",
        "Domain specialist\ntraining (3 seeds)",
        "Validation threshold\nand probability averaging",
        "Untouched test\nevaluation + bootstrap CI",
    ]
    xs = np.linspace(0.08, 0.92, len(labels))
    colors = ["#DCE8F0", "#F3E2E4", "#E5EEE8", "#DCE8F0", "#F1E8D5", "#E5EEE8"]
    for i, (x, label, color) in enumerate(zip(xs, labels, colors)):
        ax.text(x, 0.5, label, ha="center", va="center", fontsize=9.2,
                bbox=dict(boxstyle="round,pad=0.45", facecolor=color, edgecolor="#40596B", linewidth=1.1))
        if i < len(labels) - 1:
            ax.annotate("", xy=(xs[i + 1] - 0.075, 0.5), xytext=(x + 0.075, 0.5),
                        arrowprops=dict(arrowstyle="->", color="#40596B", lw=1.5))
    ax.set_title("Leakage-aware modeling and evaluation workflow", fontsize=13, fontweight="bold", pad=12)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def performance_figure(results, path):
    domains = list(DOMAIN_LABELS)
    metrics = ["mcc", "pr_auc", "balanced_accuracy"]
    labels = ["MCC", "PR-AUC", "Balanced accuracy"]
    x = np.arange(len(domains)); width = 0.23
    fig, ax = plt.subplots(figsize=(8.4, 4.6))
    palette = ["#247BA0", "#D1495B", "#2A9D6F"]
    for i, (metric, label, color) in enumerate(zip(metrics, labels, palette)):
        values = [results[d]["test"][metric] for d in domains]
        bars = ax.bar(x + (i - 1) * width, values, width, label=label, color=color)
        ax.bar_label(bars, fmt="%.3f", padding=2, fontsize=8)
    ax.set_xticks(x, [DOMAIN_LABELS[d] for d in domains])
    ax.set_ylim(0, 1.06); ax.set_ylabel("Test score")
    ax.set_title("Final three-seed ensemble performance")
    ax.grid(axis="y", alpha=0.25); ax.legend(frameon=False, ncol=3, loc="lower center")
    fig.tight_layout(); fig.savefig(path, dpi=300, bbox_inches="tight"); plt.close(fig)


def confusion_figure(results, path):
    fig, axes = plt.subplots(1, 3, figsize=(9.2, 3.15))
    for ax, domain in zip(axes, DOMAIN_LABELS):
        cm = np.array(results[domain]["test"]["confusion_matrix"])
        ax.imshow(cm, cmap="Blues", vmin=0, vmax=cm.max())
        for (i, j), value in np.ndenumerate(cm):
            ax.text(j, i, f"{value:,}", ha="center", va="center",
                    color="white" if value > cm.max() * 0.55 else "#172A3A", fontweight="bold")
        ax.set_title(DOMAIN_LABELS[domain]); ax.set_xticks([0, 1], ["Non-promoter", "Promoter"], rotation=25, ha="right")
        ax.set_yticks([0, 1], ["Non-promoter", "Promoter"] if domain == "bacteria" else ["", ""])
        ax.set_xlabel("Predicted")
    axes[0].set_ylabel("True")
    fig.suptitle("Final test confusion matrices", fontweight="bold")
    fig.tight_layout(); fig.savefig(path, dpi=300, bbox_inches="tight"); plt.close(fig)


def build_document(results, counts, figure_paths):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65); section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7); section.right_margin = Inches(0.7)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"; styles["Normal"].font.size = Pt(10)
    for name, size, color in [("Title", 22, "17324D"), ("Heading 1", 15, "17324D"), ("Heading 2", 12, "315A72")]:
        styles[name].font.name = "Arial"; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor.from_string(color)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Domain-Specialized Deep-Learning Framework for Promoter Prediction")
    run.bold = True; run.font.size = Pt(22); run.font.color.rgb = RGBColor(23, 50, 77)
    subtitle = doc.add_paragraph("Complete modeling flow and finalized fixed-configuration ensemble results")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True; subtitle.runs[0].font.size = Pt(11)
    doc.add_paragraph()

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "This report documents the implemented leakage-aware promoter classification workflow and the finalized results of fixed-configuration, three-seed domain specialists. "
        "The bacterial and archaeal systems use Multi-scale CNN ensembles; the eukaryotic system uses a Residual CNN ensemble. Ensemble probabilities are averaged across seeds, "
        "with decision thresholds chosen solely from validation predictions before one-time evaluation on untouched test partitions."
    )
    doc.add_paragraph(
        "The strongest result was obtained for Archaea (test MCC 0.842), followed by Eukaryota (0.708) and Bacteria (0.617). These results are finalized for the current fixed configurations. "
        "They are not evidence that these architectures outperform all five implemented candidates because the full architecture comparison, unified-model study, interpretability suite, and ablations remain incomplete."
    )

    doc.add_heading("1. Dataset and leakage controls", level=1)
    doc.add_paragraph(
        "Only FASTA files in data/3) finalized were used. All bacterial sequences were 81 nt, archaeal sequences 100 nt, and eukaryotic sequences 251 nt. "
        "Exact sequences and their reverse complements share a deterministic group identifier so they cannot cross partitions. Six archaeal promoter/non-promoter reverse-complement conflict pairs "
        "(12 records) were quarantined. No exact cross-label sequence overlaps were detected. One isolated N base was represented by an all-zero nucleotide vector with a separate validity mask."
    )
    dataset_rows = []
    for domain in DOMAIN_LABELS:
        promoter = sum(counts[(domain, "promoter", p)] for p in ("train", "validation", "test"))
        negative = sum(counts[(domain, "non-promoter", p)] for p in ("train", "validation", "test"))
        train = sum(counts[(domain, c, "train")] for c in ("promoter", "non-promoter"))
        validation = sum(counts[(domain, c, "validation")] for c in ("promoter", "non-promoter"))
        test = sum(counts[(domain, c, "test")] for c in ("promoter", "non-promoter"))
        dataset_rows.append([DOMAIN_LABELS[domain], f"{promoter:,}", f"{negative:,}", f"{train:,}", f"{validation:,}", f"{test:,}", f"{promoter+negative:,}"])
    doc.add_paragraph("Table 1. Leakage-controlled dataset composition", style=None).runs[0].bold = True
    add_table(doc, ["Domain", "Promoter", "Non-promoter", "Train", "Validation", "Test", "Total"], dataset_rows)

    doc.add_heading("2. Modeling flow", level=1)
    doc.add_picture(str(figure_paths[0]), width=Inches(7.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Figure 1. End-to-end data, training, threshold-selection, and evaluation workflow.").alignment = WD_ALIGN_PARAGRAPH.CENTER
    steps = [
        ("Audit and metadata", "Parse finalized FASTA files, preserve original headers, assign stable internal identifiers, compute hashes and sequence-level metadata."),
        ("Grouped splitting", "Create deterministic 70/15/15 train, validation, and test partitions stratified by domain, organism, and class while grouping exact/reverse-complement sequences."),
        ("Encoding", "Encode A/C/G/T as four-channel one-hot vectors. N and padding remain all-zero; explicit masks prevent padding from affecting recurrence, attention, or pooling."),
        ("Specialist architectures", "Use Multi-scale CNNs for Bacteria and Archaea to capture motifs at multiple widths, and a compact Residual CNN for the much larger 251-nt eukaryotic dataset."),
        ("Training", "Train each fixed architecture with seeds 2025, 2026, and 2027 using weighted BCE, organism/class-aware sampling, AdamW, gradient clipping, deterministic seeds, checkpointing, and early stopping."),
        ("Selection and ensemble", "Restore each seed's best validation-MCC checkpoint, average seed probabilities, and select the ensemble threshold by maximum validation MCC."),
        ("Final evaluation", "Apply the frozen ensemble and threshold once to the untouched test set; calculate discrimination, classification, calibration, confusion matrices, and 1,000-resample bootstrap intervals."),
    ]
    for heading, body in steps:
        paragraph = doc.add_paragraph(style=None)
        paragraph.add_run(f"{heading}: ").bold = True
        paragraph.add_run(body)

    doc.add_heading("3. Finalized ensemble results", level=1)
    result_rows = []
    for domain in DOMAIN_LABELS:
        item = results[domain]; m = item["test"]
        architecture = "Multi-scale CNN" if item["model"] == "multiscale_cnn" else "Residual CNN"
        result_rows.append([DOMAIN_LABELS[domain], architecture, f"{item['threshold']:.3f}", f"{m['accuracy']:.3f}", f"{m['sensitivity']:.3f}", f"{m['specificity']:.3f}", f"{m['mcc']:.3f}", f"{m['pr_auc']:.3f}", f"{m['roc_auc']:.3f}"])
    doc.add_paragraph("Table 2. Final untouched-test performance", style=None).runs[0].bold = True
    add_table(doc, ["Domain", "Architecture", "Threshold", "Accuracy", "Sensitivity", "Specificity", "MCC", "PR-AUC", "ROC-AUC"], result_rows)
    doc.add_picture(str(figure_paths[1]), width=Inches(6.9)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Figure 2. Principal test metrics for the finalized three-seed ensembles.").alignment = WD_ALIGN_PARAGRAPH.CENTER

    ci_rows = []
    for domain in DOMAIN_LABELS:
        ci = results[domain]["test_bootstrap_95_ci"]
        ci_rows.append([DOMAIN_LABELS[domain], f"{ci['mcc'][0]:.3f}-{ci['mcc'][1]:.3f}", f"{ci['pr_auc'][0]:.3f}-{ci['pr_auc'][1]:.3f}", f"{ci['roc_auc'][0]:.3f}-{ci['roc_auc'][1]:.3f}"])
    doc.add_paragraph("Table 3. Bootstrap 95% confidence intervals (1,000 resamples)", style=None).runs[0].bold = True
    add_table(doc, ["Domain", "MCC 95% CI", "PR-AUC 95% CI", "ROC-AUC 95% CI"], ci_rows)
    doc.add_picture(str(figure_paths[2]), width=Inches(7.0)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Figure 3. Confusion matrices using validation-selected ensemble thresholds.").alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("4. Interpretation of finalized results", level=1)
    doc.add_paragraph(
        "Archaea achieved the highest discrimination and classification balance (MCC 0.842; PR-AUC 0.954), with both false-positive and false-negative rates controlled. "
        "Eukaryota achieved high specificity and precision but lower sensitivity, indicating a conservative promoter decision boundary. Bacterial performance was weaker and less certain, "
        "consistent with a smaller, organism-heterogeneous dataset and greater validation-to-test variation."
    )
    doc.add_paragraph(
        "Relative to the corresponding single-seed models, the three-seed ensembles improved test MCC by 0.056 for Bacteria, 0.032 for Archaea, and 0.013 for Eukaryota. "
        "The consistent direction of change supports variance reduction through probability averaging, although paired statistical testing is still needed before claiming significance."
    )

    doc.add_heading("5. Reproducibility and saved artifacts", level=1)
    for text in [
        "Data manifest: data/splits/split_manifest.csv",
        "Per-seed checkpoints and predictions: results/checkpoints/<domain>/<model>/<config>/seed_<seed>/",
        "Ensemble metrics: results/main_tables/multiseed_ensemble.json",
        "Ensemble predictions: results/predictions/<domain>_multiseed_ensemble_test.csv",
        "Random seeds: 2025, 2026, and 2027",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("6. Publication-readiness statement", level=1)
    doc.add_paragraph(
        "The dataset audit, leakage-controlled split, fixed-configuration specialist training, three-seed ensembling, untouched-test evaluation, and bootstrap uncertainty estimates are complete and reproducible. "
        "These results can be reported as finalized specialist-ensemble experiments. The broader proposed publication is not yet complete because it still requires the planned five-architecture comparison, "
        "unified domain-aware model, group-aware cross-validation, biological controls, interpretability analyses, ablations, paired statistical tests, and final model exports."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Promoter Prediction Research Project | Reproducible fixed-configuration results")
    doc.save(OUTPUT)


def main():
    REPORTS.mkdir(parents=True, exist_ok=True); FIGURES.mkdir(parents=True, exist_ok=True)
    results, counts = load_data()
    figure_paths = [FIGURES / "modeling_workflow.png", FIGURES / "ensemble_performance.png", FIGURES / "ensemble_confusion_matrices.png"]
    workflow_figure(figure_paths[0]); performance_figure(results, figure_paths[1]); confusion_figure(results, figure_paths[2])
    build_document(results, counts, figure_paths)
    print(OUTPUT)


if __name__ == "__main__":
    main()
