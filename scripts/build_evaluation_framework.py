from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "Evaluation Framework.docx"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def add_formula(document, formula):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(formula)
    run.bold = True
    run.font.name = "Cambria Math"
    run.font.size = Pt(11.5)


def add_metric(document, number, name, formula, explanation):
    document.add_heading(f"{number} {name}", level=2)
    add_formula(document, formula)
    document.add_paragraph(explanation)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        shade(cell, "24435C")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            if row_index % 2:
                shade(cells[index], "EAF0F4")
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(9)
    return table


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(11)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12.5)):
        document.styles[name].font.name = "Times New Roman"
        document.styles[name].font.size = Pt(size)
        document.styles[name].font.color.rgb = RGBColor(26, 55, 76)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Evaluation Framework")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(26, 55, 76)
    subtitle = document.add_paragraph("Performance metrics, decision rules, and uncertainty estimation")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True

    document.add_heading("1. Evaluation overview", level=1)
    document.add_paragraph(
        "Promoter prediction was evaluated as a binary classification problem, where the positive class represented promoter sequences and the negative class represented non-promoter sequences. "
        "Evaluation was conducted using validation-selected decision thresholds and untouched test predictions. Matthews correlation coefficient (MCC) was treated as the primary classification metric, "
        "while precision-recall area under the curve (PR-AUC), balanced accuracy, sensitivity, and specificity were treated as principal secondary measures. Accuracy was reported for completeness but was not used as the sole indicator of performance."
    )

    document.add_heading("2. Confusion-matrix definitions", level=1)
    document.add_paragraph(
        "Let TP, TN, FP, and FN denote the numbers of true-positive, true-negative, false-positive, and false-negative predictions, respectively. These quantities were defined as follows."
    )
    add_table(
        document,
        ["Term", "Definition", "Promoter-prediction interpretation"],
        [
            ["True positive (TP)", "Positive observation predicted as positive", "A promoter correctly classified as a promoter"],
            ["True negative (TN)", "Negative observation predicted as negative", "A non-promoter correctly classified as a non-promoter"],
            ["False positive (FP)", "Negative observation predicted as positive", "A non-promoter incorrectly classified as a promoter"],
            ["False negative (FN)", "Positive observation predicted as negative", "A promoter incorrectly classified as a non-promoter"],
        ],
    )
    document.add_paragraph("The binary confusion matrix was organized as:")
    add_formula(document, "Confusion matrix = [[TN, FP], [FN, TP]]")

    document.add_heading("3. Threshold-dependent classification metrics", level=1)
    add_metric(
        document, "3.1", "Accuracy",
        "Accuracy = (TP + TN) / (TP + TN + FP + FN)",
        "Accuracy is the proportion of all test sequences that were classified correctly. Although intuitive, it can be misleading when promoter and non-promoter classes are imbalanced."
    )
    add_metric(
        document, "3.2", "Sensitivity (Recall or True-Positive Rate)",
        "Sensitivity = Recall = TPR = TP / (TP + FN)",
        "Sensitivity measures the proportion of true promoters detected by the model. A low sensitivity indicates that many promoters were missed."
    )
    add_metric(
        document, "3.3", "Specificity (True-Negative Rate)",
        "Specificity = TNR = TN / (TN + FP)",
        "Specificity measures the proportion of non-promoters correctly rejected. A low specificity indicates that many non-promoters were incorrectly labeled as promoters."
    )
    add_metric(
        document, "3.4", "Precision (Positive Predictive Value)",
        "Precision = PPV = TP / (TP + FP)",
        "Precision is the proportion of predicted promoters that were genuine promoters. It quantifies the reliability of a positive prediction."
    )
    add_metric(
        document, "3.5", "F1-score",
        "F1 = 2 × (Precision × Recall) / (Precision + Recall) = 2TP / (2TP + FP + FN)",
        "The F1-score is the harmonic mean of precision and recall. It emphasizes positive-class performance but does not explicitly incorporate true negatives."
    )
    add_metric(
        document, "3.6", "Matthews Correlation Coefficient",
        "MCC = (TP × TN − FP × FN) / √[(TP + FP)(TP + FN)(TN + FP)(TN + FN)]",
        "MCC summarizes all four confusion-matrix components and ranges from −1 to +1. A value of +1 represents perfect classification, 0 indicates performance no better than random association, and −1 represents complete disagreement. "
        "MCC was the primary metric because it remains informative when the two classes differ in size. If the denominator was zero, the implementation returned a neutral value according to the metric library's defined behavior."
    )
    add_metric(
        document, "3.7", "Balanced Accuracy",
        "Balanced accuracy = (Sensitivity + Specificity) / 2",
        "Balanced accuracy assigns equal importance to promoter detection and non-promoter rejection, regardless of class prevalence."
    )
    add_metric(
        document, "3.8", "False-Positive Rate",
        "FPR = FP / (FP + TN) = 1 − Specificity",
        "The false-positive rate is the proportion of non-promoters incorrectly classified as promoters."
    )
    add_metric(
        document, "3.9", "False-Negative Rate",
        "FNR = FN / (FN + TP) = 1 − Sensitivity",
        "The false-negative rate is the proportion of promoters incorrectly classified as non-promoters."
    )

    document.add_heading("4. Threshold-independent ranking metrics", level=1)
    add_metric(
        document, "4.1", "Receiver Operating Characteristic Curve",
        "ROC(t) = [FPR(t), TPR(t)] for decision thresholds t",
        "The receiver operating characteristic curve plots sensitivity against the false-positive rate as the decision threshold is varied. Each point describes the trade-off between detecting promoters and incorrectly labeling non-promoters."
    )
    add_metric(
        document, "4.2", "Area Under the ROC Curve",
        "ROC-AUC = ∫₀¹ TPR(FPR⁻¹(u)) du",
        "ROC-AUC measures the probability that a randomly selected promoter receives a higher score than a randomly selected non-promoter. A value of 0.5 represents random ranking and 1.0 represents perfect discrimination. "
        "The area was calculated numerically from the empirical ROC curve."
    )
    add_metric(
        document, "4.3", "Precision-Recall Curve",
        "PR(t) = [Recall(t), Precision(t)] for decision thresholds t",
        "The precision-recall curve displays the relationship between promoter recall and the precision of positive predictions as the threshold changes. It is especially informative when the positive class is imbalanced."
    )
    add_metric(
        document, "4.4", "Area Under the Precision-Recall Curve",
        "PR-AUC = Σₙ (Recallₙ − Recallₙ₋₁) × Precisionₙ",
        "PR-AUC summarizes promoter-focused ranking performance across thresholds. The implementation used average precision, which forms a weighted sum of precision values using the increase in recall between successive thresholds. "
        "A higher value indicates that promoters can be recovered with fewer false-positive predictions."
    )

    document.add_heading("5. Probability-quality metric", level=1)
    add_metric(
        document, "5.1", "Brier Score",
        "Brier score = (1/N) Σᵢ₌₁ᴺ (pᵢ − yᵢ)²",
        "Here, pᵢ is the predicted promoter probability and yᵢ is the binary class label for sequence i. The Brier score measures the mean squared error of probabilistic predictions and ranges from 0 to 1 for a binary outcome. "
        "Lower values indicate better probability accuracy. Unlike discrimination metrics, the Brier score is sensitive to both probability calibration and classification uncertainty."
    )

    document.add_heading("6. Validation-based threshold selection", level=1)
    document.add_paragraph(
        "The class-decision threshold was selected using validation predictions only. Let pᵢ denote the validation promoter probability for observation i. Every distinct validation probability was considered as a candidate threshold, and the threshold producing the maximum validation MCC was selected:"
    )
    add_formula(document, "t* = arg maxₜ MCC(yvalidation, I[pvalidation ≥ t])")
    document.add_paragraph(
        "where I[·] is the indicator function. An exact sorted cumulative implementation was used to evaluate all distinct thresholds efficiently. If multiple thresholds produced the same maximum MCC, the threshold closest to 0.50 was selected. "
        "The resulting threshold was frozen before evaluation on the untouched test set."
    )

    document.add_heading("7. Multi-seed ensemble evaluation", level=1)
    document.add_paragraph(
        "For each domain, predictions from the three independently initialized models were combined through arithmetic mean probability averaging. For a sequence x and S = 3 seed models, the ensemble probability was:"
    )
    add_formula(document, "pensemble(x) = (1/S) Σₛ₌₁ˢ pₛ(x), where S = 3")
    document.add_paragraph(
        "The ensemble threshold was selected from averaged validation probabilities using the same maximum-MCC rule. Test probabilities were averaged only after the ensemble procedure and threshold had been fixed. No ensemble weight or threshold was fitted using test predictions."
    )

    document.add_heading("8. Bootstrap confidence intervals", level=1)
    document.add_paragraph(
        "Uncertainty in final test performance was estimated using non-parametric bootstrap resampling. For a test set containing N observations, each bootstrap replicate sampled N observation indices with replacement from the original test predictions. "
        "MCC, ROC-AUC, and PR-AUC were recalculated for each resample using the frozen ensemble threshold where applicable. A total of B = 1,000 bootstrap replicates were generated with random seed 2025."
    )
    add_formula(document, "95% CIθ = [Q₀.₀₂₅({θ̂ᵦ}ᵦ₌₁ᴮ), Q₀.₉₇₅({θ̂ᵦ}ᵦ₌₁ᴮ)]")
    document.add_paragraph(
        "In this expression, θ̂ᵦ is the metric calculated in bootstrap replicate b, and Q denotes the empirical quantile. The 2.5th and 97.5th percentiles formed the percentile-based 95% confidence interval. "
        "Bootstrap replicates for which a ranking metric was undefined because only one class was sampled were excluded for that metric."
    )

    document.add_heading("9. Metric-selection hierarchy", level=1)
    document.add_paragraph(
        "The following hierarchy was used when interpreting model performance:"
    )
    for item in [
        "MCC as the primary threshold-dependent measure of overall classification quality.",
        "PR-AUC as the primary threshold-independent measure emphasizing promoter retrieval.",
        "Balanced accuracy to assess equal performance across promoter and non-promoter classes.",
        "Sensitivity and specificity to examine the balance between missed promoters and false promoter calls.",
        "ROC-AUC to characterize global ranking discrimination.",
        "Brier score to assess the accuracy of predicted probabilities.",
        "Accuracy, precision, F1-score, FPR, FNR, and the confusion matrix as complementary descriptive measures.",
    ]:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("10. Reporting safeguards", level=1)
    document.add_paragraph(
        "All thresholds were selected from validation predictions, and reported test metrics were calculated only after the model checkpoint, ensemble rule, and threshold had been fixed. Metrics were calculated on identical domain-specific test samples to preserve comparability. "
        "The test set was not used for training, early stopping, architecture selection, sampling-weight calculation, class-weight calculation, or threshold optimization. Accuracy was never interpreted in isolation, and numerical differences were reported together with bootstrap uncertainty where available."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Evaluation Framework | Domain-Specialized Promoter Prediction")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
