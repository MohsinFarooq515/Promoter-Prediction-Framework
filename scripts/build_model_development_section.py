from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "Model_Development_Section.docx"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        shade(cell, "24435C")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8.5)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            if row_index % 2:
                shade(cells[index], "EAF0F4")
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(8.5)
    return table


def add_labeled_paragraph(document, label, text):
    paragraph = document.add_paragraph()
    paragraph.add_run(label).bold = True
    paragraph.add_run(text)
    return paragraph


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(11)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)):
        document.styles[name].font.name = "Times New Roman"
        document.styles[name].font.size = Pt(size)
        document.styles[name].font.color.rgb = RGBColor(26, 55, 76)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Model Development")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(26, 55, 76)
    subtitle = document.add_paragraph("Thesis-ready methodology section")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True

    document.add_heading("1. Model development strategy", level=1)
    document.add_paragraph(
        "The model-development stage was designed to construct a leakage-aware deep-learning framework for promoter classification across Bacteria, Archaea, and Eukaryota. "
        "The problem was formulated as binary sequence classification, where each nucleotide sequence was assigned either a promoter or non-promoter label. Because promoter organization, "
        "motif composition, and sequence length differ substantially among biological domains, the modeling strategy did not assume that a single architecture would be optimal for every domain. "
        "Instead, five candidate neural architectures were implemented within a common training and evaluation framework, while domain-specific specialist models were trained using the natural input length of each domain."
    )
    document.add_paragraph(
        "The five implemented architectures were: (i) a one-hot Residual Convolutional Neural Network (Residual CNN), (ii) a Multi-scale CNN, (iii) a CNN followed by a bidirectional "
        "Long Short-Term Memory network (CNN-BiLSTM), (iv) a CNN-Transformer, and (v) a proposed Domain-aware Multi-scale CNN-Transformer. These architectures were selected to represent "
        "complementary mechanisms for learning local motifs, multi-resolution sequence patterns, long-range dependencies, and shared or domain-specific promoter representations. Classical machine-learning "
        "algorithms were deliberately excluded so that the comparison remained focused on deep neural sequence models."
    )

    document.add_heading("2. Input representation and masking", level=1)
    document.add_paragraph(
        "Each nucleotide sequence was converted into a four-channel one-hot matrix. Adenine, cytosine, guanine, and thymine were represented as [1,0,0,0], [0,1,0,0], [0,0,1,0], and [0,0,0,1], "
        "respectively. An ambiguous N base was encoded as [0,0,0,0] rather than being replaced with a fabricated nucleotide. Two masks were retained: a sequence-position mask identified positions belonging "
        "to the original sequence, whereas a nucleotide-validity mask distinguished canonical nucleotides from ambiguous positions. This separation ensured that an N base remained part of the biological sequence "
        "without being interpreted as a fifth nucleotide category."
    )
    document.add_paragraph(
        "The finalized bacterial, archaeal, and eukaryotic sequences had natural lengths of 81, 100, and 251 nucleotides, respectively. Domain-specific models therefore operated directly at these lengths and did not "
        "pad all domains to a common global maximum. The implementation nevertheless supported right-padding for shorter valid inputs. Padded positions were encoded as zero vectors and excluded from recurrent processing, "
        "Transformer attention, global pooling, prediction, and interpretation through explicit Boolean masks. This design prevented sequence length or padding from becoming an artificial class indicator."
    )

    document.add_heading("3. Candidate neural architectures", level=1)
    document.add_heading("3.1 One-hot Residual CNN", level=2)
    document.add_paragraph(
        "The Residual CNN was developed as a strong convolutional baseline. The one-hot input was transposed to a channel-first representation and passed through a one-dimensional convolutional projection. "
        "The projected features were processed by residual blocks containing Conv1D, batch normalization, Gaussian Error Linear Unit activation, dropout, a second Conv1D operation, and batch normalization. "
        "The input to each block was added to its transformed output before activation, allowing gradients to propagate through deeper convolutional stacks and reducing degradation during optimization. "
        "The final sequence representation combined masked global-average and masked global-maximum pooling. Their concatenation was passed through dropout and a linear binary-classification head."
    )

    document.add_heading("3.2 Multi-scale CNN", level=2)
    document.add_paragraph(
        "The Multi-scale CNN was designed to detect promoter signals occurring at different motif widths. The same one-hot sequence was processed in parallel convolutional branches with kernels of 3, 5, 7, and 11 nucleotides in the standard screening configuration. "
        "Each branch contained convolution, batch normalization, GELU activation, residual processing, and dropout. Branch outputs were concatenated at each position and projected into a shared hidden space. "
        "Masked mean and maximum pooling then summarized both widespread compositional evidence and the strongest motif activation. The fused representation was supplied to a dropout-regularized linear classifier. "
        "This model was biologically appropriate for bacterial and archaeal promoters because it could jointly learn short core elements and broader contextual or spacing-dependent patterns."
    )

    document.add_heading("3.3 CNN-BiLSTM", level=2)
    document.add_paragraph(
        "The CNN-BiLSTM combined local motif extraction with bidirectional sequential modeling. A Conv1D feature extractor first transformed the four-channel nucleotide matrix into a higher-dimensional feature sequence. "
        "The feature sequence was then packed according to the unpadded sequence lengths and passed through a bidirectional LSTM. Packing ensured that padded positions did not update recurrent states. "
        "Forward and reverse hidden representations were combined through masked global-average and maximum pooling, followed by dropout and a linear promoter-classification head. "
        "This architecture was intended to represent dependencies in both upstream-to-downstream and downstream-to-upstream directions."
    )

    document.add_heading("3.4 CNN-Transformer", level=2)
    document.add_paragraph(
        "The CNN-Transformer used a convolutional front end to extract motif-level features before global contextual modeling. Conv1D features were augmented with sinusoidal positional encodings and supplied to a pre-normalized Transformer encoder. "
        "Multi-head self-attention enabled every valid position to interact with other sequence positions, while the key-padding mask prevented attention to padded positions. The encoder output was summarized using masked mean and maximum pooling, "
        "and the resulting representation was passed to a dropout-regularized classifier. The Transformer configuration supported multiple attention heads, encoder layers, feed-forward expansion factors, and dropout rates, subject to hidden-dimension/head compatibility checks."
    )

    document.add_heading("3.5 Proposed Domain-aware Multi-scale CNN-Transformer", level=2)
    document.add_paragraph(
        "The proposed model was implemented to support unified learning across all three domains. Separate bacterial, archaeal, and eukaryotic input encoders projected the four-channel inputs into a common latent dimension. "
        "The encoded features were analyzed through parallel multi-scale convolutional branches and fused by linear projection. A learned domain embedding was added to the fused sequence representation before positional encoding and shared Transformer processing. "
        "Masked pooling produced a shared biological representation used by three output components: a shared promoter head, the corresponding domain-specific promoter head, and an auxiliary domain-classification head. "
        "The promoter logit could be formed as a validation-weighted combination of the shared and domain-specific promoter outputs. The auxiliary task was assigned a smaller configurable weight so that promoter prediction remained the primary objective."
    )

    document.add_paragraph("Table 1. Functional roles of the implemented candidate architectures").runs[0].bold = True
    add_table(
        document,
        ["Architecture", "Primary modeling capability", "Mask handling", "Study role"],
        [
            ["Residual CNN", "Deep local motif extraction", "Masked pooling", "Strong convolutional baseline; finalized Eukaryota specialist"],
            ["Multi-scale CNN", "Motifs at multiple sequence widths", "Masked pooling", "Finalized Bacteria and Archaea specialists"],
            ["CNN-BiLSTM", "Local motifs and bidirectional order", "Packed recurrence and masked pooling", "Implemented comparison candidate"],
            ["CNN-Transformer", "Local motifs and global dependencies", "Attention padding mask and masked pooling", "Implemented comparison candidate"],
            ["Domain-aware model", "Shared and domain-specific promoter grammar", "Masked attention and pooling", "Implemented proposed unified architecture"],
        ],
    )

    document.add_heading("4. Loss function and sampling strategy", level=1)
    document.add_paragraph(
        "Binary cross-entropy with logits was used as the primary promoter-classification loss. To reduce the effect of class imbalance, the positive-class weight was calculated exclusively from the training partition as the ratio of training non-promoters to training promoters. "
        "No validation or test labels were used to derive loss weights. The implementation also supported unweighted binary cross-entropy and focal loss as alternative configurations, although the finalized fixed-configuration specialist runs used weighted binary cross-entropy."
    )
    document.add_paragraph(
        "Training examples were sampled with inverse-frequency weights defined over the combination of biological domain, organism, and class. Consequently, smaller organism/class strata received greater sampling probability without permanently deleting majority-class sequences. "
        "Sampling was applied only to the training set. Validation and test partitions retained their natural distributions. For the proposed multi-task model, the total objective combined shared promoter loss, domain-specific promoter loss, and a lower-weight domain-classification loss."
    )

    document.add_heading("5. Optimization and regularization", level=1)
    document.add_paragraph(
        "Models were optimized on CPU using AdamW with a learning rate of 3 × 10^-4 and weight decay of 1 × 10^-4 in the fixed screening configuration. Gradient norms were clipped to 1.0 to reduce unstable updates. "
        "Dropout was applied within feature-processing blocks and before classification. Batch normalization was used in convolutional modules, while Transformer layers employed pre-normalization. "
        "The tanh approximation of GELU was used because it provided numerically stable gradients in the available PyTorch CPU environment."
    )
    document.add_paragraph(
        "A ReduceLROnPlateau scheduler monitored validation MCC and reduced the learning rate when progress stalled. Early stopping used validation MCC with domain-appropriate patience, while a maximum epoch budget bounded every run. "
        "The test set was not consulted during optimization or checkpoint selection. Each epoch saved the current model, optimizer, scheduler, best validation MCC, early-stopping counter, configuration, and complete training history. "
        "Separate checkpoints were retained for the highest validation MCC and highest validation PR-AUC, and interrupted runs resumed from the latest checkpoint without resetting early-stopping state."
    )

    document.add_heading("6. Validation-based model selection", level=1)
    document.add_paragraph(
        "Matthews correlation coefficient was selected as the primary model-development criterion because it incorporates all four cells of the confusion matrix and remains informative under class imbalance. "
        "PR-AUC and balanced accuracy were treated as important secondary criteria. For each trained seed, the checkpoint with the highest validation MCC was restored. The operating threshold was not fixed at 0.50; instead, "
        "candidate thresholds were evaluated on validation probabilities and the threshold maximizing validation MCC was selected using an exact sorted cumulative calculation. The frozen threshold was subsequently applied to the corresponding untouched test predictions."
    )

    document.add_heading("7. Multi-seed training and ensemble construction", level=1)
    document.add_paragraph(
        "To improve robustness without conducting hyperparameter optimization, the selected specialist configurations were trained independently with random seeds 2025, 2026, and 2027. "
        "The bacterial and archaeal systems used the Multi-scale CNN, whereas the eukaryotic system used the Residual CNN. These choices reflected the observed suitability of multi-resolution motif extraction for the shorter prokaryotic/archaeal windows and the computational efficiency of residual convolution for the much larger 251-nt eukaryotic dataset."
    )
    document.add_paragraph(
        "For each domain, promoter probabilities from the three frozen seed models were averaged arithmetically for every validation sample. A single ensemble threshold was then selected by maximizing MCC on the averaged validation probabilities. "
        "After freezing this threshold, probabilities were averaged for each test sample and converted to class predictions. This approach reduced sensitivity to random initialization without learning ensemble weights from the test set. "
        "Bootstrap confidence intervals were calculated from 1,000 test-set resamples for MCC, ROC-AUC, and PR-AUC."
    )

    document.add_heading("8. Final model configurations", level=1)
    document.add_paragraph("Table 2. Final fixed-configuration specialist ensembles").runs[0].bold = True
    add_table(
        document,
        ["Domain", "Architecture", "Input length", "Seeds", "Ensemble rule", "Validation-selected threshold"],
        [
            ["Bacteria", "Multi-scale CNN", "81 nt", "2025, 2026, 2027", "Mean probability", "0.876"],
            ["Archaea", "Multi-scale CNN", "100 nt", "2025, 2026, 2027", "Mean probability", "0.598"],
            ["Eukaryota", "Residual CNN", "251 nt", "2025, 2026, 2027", "Mean probability", "0.560"],
        ],
    )
    document.add_paragraph(
        "The finalized bacterial and archaeal Multi-scale CNNs used four convolutional motif scales (3, 5, 7, and 11 nt), residual processing, masked dual pooling, and a shared hidden projection. "
        "The finalized eukaryotic Residual CNN used a compact channel width and two residual blocks with a larger batch size to support full-dataset CPU training. All configurations used the complete training partition; no permanent undersampling or dataset-size reduction was performed."
    )

    document.add_heading("9. Reproducibility safeguards", level=1)
    for item in [
        "Stable internal sequence identifiers and preserved original FASTA headers.",
        "Deterministic group-aware partitions preventing exact-sequence and reverse-complement leakage.",
        "Training-only calculation of class weights and sampling probabilities.",
        "Deterministic random seeds and CPU thread control.",
        "Per-epoch checkpoints containing model and optimizer state.",
        "Validation-only checkpoint and threshold selection.",
        "Saved per-seed validation/test predictions and ensemble predictions.",
        "Recorded configurations, parameter counts, durations, histories, and dataset hashes.",
    ]:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("10. Scope of the completed model development", level=1)
    document.add_paragraph(
        "All five neural architectures were implemented and verified for valid output dimensions and mask-aware processing. The completed empirical results, however, correspond to the fixed-configuration, three-seed specialist ensembles described above. "
        "The full five-architecture comparison for every domain, focused hyperparameter optimization, unified-model evaluation, group-aware cross-validation, biological controls, interpretability analyses, and proposed-model ablations were not completed. "
        "Accordingly, the finalized specialist models can be reported as reproducible trained systems, but they should not be described as globally optimal among all implemented architectures until the remaining paired comparisons are performed."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Model Development | Domain-Specialized Promoter Prediction")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
