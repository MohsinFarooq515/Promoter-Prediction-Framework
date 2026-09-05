from pathlib import Path
import json
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Detailed_Results_Chapter.docx"
DOMAINS = ("bacteria", "archaea", "eukaryota")
LABEL = {"bacteria": "Bacteria", "archaea": "Archaea", "eukaryota": "Eukaryota"}


def load(path):
    return json.loads((ROOT / path).read_text(encoding="utf-8"))


def fmt(x, n=3):
    return f"{x:.{n}f}"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def table(doc, headers, rows, size=8.1):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    repeat_header(t.rows[0])
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(c, "1F4E78")
        for r in c.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(size)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2:
                shade(cells[i], "EAF2F8")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(size)
    t.autofit = True
    return t


def body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.3)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(6)
    return p


def caption(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.runs[0].bold = True


def picture(doc, path, width, text):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(text)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True


def configure(doc):
    s = doc.sections[0]
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.9)
    s.right_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(31, 78, 121)
        st.paragraph_format.keep_with_next = True
    foot = s.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run("Results | Domain-Aware Promoter Prediction")


def build():
    ensemble = load("results/main_tables/multiseed_ensemble.json")
    hybrid = {d: load(f"results/hybrid_models/{d}/metadata.json") for d in DOMAINS}
    ablation = load("results/main_tables/feature_ablation_oof.json")

    doc = Document()
    configure(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CHAPTER 4\nRESULTS")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(31, 78, 121)

    body(doc, "This chapter presents the empirical findings obtained from the leakage-controlled promoter-classification experiments. Results are reported in three layers. First, the performance of the finalized three-seed neural ensembles is examined for Bacteria, Archaea, and Eukaryota. Second, organism-specific feature-ablation experiments are used to determine which explicit sequence descriptors contributed useful complementary information. Third, the final hybrid models, which combine neural-network and Random Forest probabilities, are compared with the corresponding CNN-only ensembles. All operating thresholds were fixed from development or validation predictions before the untouched test data were evaluated. Consequently, the test results reported here represent confirmatory performance estimates rather than values selected to maximize test-set accuracy.")

    doc.add_heading("4.1 Overview of completed experiments", level=1)
    body(doc, "The finalized neural experiments used Multi-scale convolutional neural networks for the bacterial and archaeal domains and a Residual CNN for Eukaryota. Each architecture was independently trained with random seeds 2025, 2026, and 2027. For each test sequence, the three promoter probabilities were averaged and converted to a binary prediction using a threshold selected from the corresponding averaged validation probabilities. The domain test sets contained 1,169 bacterial sequences, 1,589 archaeal sequences, and 12,100 eukaryotic sequences. The completed neural evaluation therefore covered 14,858 non-quarantined test records.")
    body(doc, "The three-seed neural ensembles were then supplemented by Random Forest models trained on explicit k-mer, positional, and thermodynamic descriptors. The final bacterial and archaeal forests contained 500 trees, whereas the saved eukaryotic hybrid used 300 trees. Hybrid weights were selected from validation predictions. The resulting models were not constrained to give equal importance to learned and engineered representations: the bacterial hybrid assigned 90% of its probability to the CNN, while the archaeal and eukaryotic hybrids assigned 45% to the CNN and 55% to the Random Forest. These weights provide direct evidence about the relative complementarity of the two representation families in each domain.")

    doc.add_heading("4.2 Performance of the three-seed neural ensembles", level=1)
    caption(doc, "Table 4.1. Test performance of finalized three-seed neural ensembles")
    rows = []
    for d in DOMAINS:
        x = ensemble[d]
        m = x["test"]
        rows.append([LABEL[d], "Multi-scale CNN" if d != "eukaryota" else "Residual CNN",
                     f"{len(x['seeds'])}", fmt(x["threshold"]), fmt(m["accuracy"]), fmt(m["sensitivity"]),
                     fmt(m["specificity"]), fmt(m["precision"]), fmt(m["f1"]), fmt(m["mcc"]),
                     fmt(m["balanced_accuracy"]), fmt(m["roc_auc"]), fmt(m["pr_auc"])])
    table(doc, ["Domain", "Architecture", "Seeds", "Threshold", "Accuracy", "Sensitivity", "Specificity", "Precision", "F1", "MCC", "Bal. acc.", "ROC-AUC", "PR-AUC"], rows, 7.0)
    picture(doc, ROOT / "results/main_figures/ensemble_performance.png", 6.8,
            "Figure 4.1. Principal test metrics of the finalized three-seed neural ensembles.")

    doc.add_heading("4.2.1 Bacterial neural ensemble", level=2)
    m = ensemble["bacteria"]["test"]
    body(doc, f"The bacterial Multi-scale CNN ensemble achieved an accuracy of {fmt(m['accuracy'])}, MCC of {fmt(m['mcc'])}, balanced accuracy of {fmt(m['balanced_accuracy'])}, ROC-AUC of {fmt(m['roc_auc'])}, and PR-AUC of {fmt(m['pr_auc'])}. The validation-selected threshold was {fmt(ensemble['bacteria']['threshold'])}, substantially higher than the conventional 0.50 cut-off. At this operating point, specificity ({fmt(m['specificity'])}) and precision ({fmt(m['precision'])}) exceeded sensitivity ({fmt(m['sensitivity'])}). The classifier was therefore conservative when assigning the promoter class: a positive bacterial prediction was relatively reliable, but approximately {100*m['false_negative_rate']:.1f}% of true promoters were missed. This asymmetry is important because accuracy alone would conceal the difference between promoter recovery and non-promoter rejection.")
    ci = ensemble["bacteria"]["test_bootstrap_95_ci"]
    body(doc, f"Bootstrap analysis produced a 95% interval of {fmt(ci['mcc'][0])}–{fmt(ci['mcc'][1])} for MCC, {fmt(ci['roc_auc'][0])}–{fmt(ci['roc_auc'][1])} for ROC-AUC, and {fmt(ci['pr_auc'][0])}–{fmt(ci['pr_auc'][1])} for PR-AUC. These intervals show that bacterial classification was clearly better than chance, while also displaying greater uncertainty than the larger eukaryotic evaluation. The bacterial domain combined E. coli and B. subtilis, and the lower sensitivity may reflect both the comparatively limited B. subtilis sample size and differences between Gram-negative and Gram-positive promoter organization.")

    doc.add_heading("4.2.2 Archaeal neural ensemble", level=2)
    m = ensemble["archaea"]["test"]
    body(doc, f"The archaeal Multi-scale CNN produced the strongest neural result. Test accuracy reached {fmt(m['accuracy'])}, MCC was {fmt(m['mcc'])}, and balanced accuracy was {fmt(m['balanced_accuracy'])}. Ranking performance was also high, with ROC-AUC of {fmt(m['roc_auc'])} and PR-AUC of {fmt(m['pr_auc'])}. Sensitivity ({fmt(m['sensitivity'])}) and specificity ({fmt(m['specificity'])}) were both strong, although the latter remained higher. Precision of {fmt(m['precision'])} indicated that most predicted archaeal promoters were correct. In contrast with the bacterial system, the archaeal threshold of {fmt(ensemble['archaea']['threshold'])} was closer to 0.50 and supported a more balanced trade-off between the two classes.")
    ci = ensemble["archaea"]["test_bootstrap_95_ci"]
    body(doc, f"The archaeal MCC interval was {fmt(ci['mcc'][0])}–{fmt(ci['mcc'][1])}; the ROC-AUC interval was {fmt(ci['roc_auc'][0])}–{fmt(ci['roc_auc'][1])}; and the PR-AUC interval was {fmt(ci['pr_auc'][0])}–{fmt(ci['pr_auc'][1])}. The lower bound of the archaeal MCC interval exceeded the upper bound of the bacterial interval in these independent bootstrap summaries. This observation supports a substantial empirical performance difference under the present datasets, although it should not be interpreted as a formal paired significance test because the domains contain different sequences.")

    doc.add_heading("4.2.3 Eukaryotic neural ensemble", level=2)
    m = ensemble["eukaryota"]["test"]
    body(doc, f"The eukaryotic Residual CNN ensemble achieved accuracy of {fmt(m['accuracy'])}, MCC of {fmt(m['mcc'])}, balanced accuracy of {fmt(m['balanced_accuracy'])}, ROC-AUC of {fmt(m['roc_auc'])}, and PR-AUC of {fmt(m['pr_auc'])}. Specificity ({fmt(m['specificity'])}) and precision ({fmt(m['precision'])}) were high, whereas sensitivity was {fmt(m['sensitivity'])}. The false-negative rate was therefore {100*m['false_negative_rate']:.1f}%, compared with a false-positive rate of only {100*m['false_positive_rate']:.1f}%. The model was effective at avoiding false promoter calls but did not recover all true promoters, which is consistent with the diversity of promoter structures represented by human, mouse, and Arabidopsis sequences.")
    ci = ensemble["eukaryota"]["test_bootstrap_95_ci"]
    body(doc, f"Because the eukaryotic test set was substantially larger, its bootstrap intervals were narrow: {fmt(ci['mcc'][0])}–{fmt(ci['mcc'][1])} for MCC, {fmt(ci['roc_auc'][0])}–{fmt(ci['roc_auc'][1])} for ROC-AUC, and {fmt(ci['pr_auc'][0])}–{fmt(ci['pr_auc'][1])} for PR-AUC. The high PR-AUC demonstrates that eukaryotic promoters generally received higher scores than negatives even though the selected classification threshold emphasized specificity. Thus, threshold-dependent sensitivity and threshold-independent ranking should be interpreted together.")

    doc.add_heading("4.3 Confusion-matrix analysis", level=1)
    picture(doc, ROOT / "results/main_figures/ensemble_confusion_matrices.png", 6.9,
            "Figure 4.2. Confusion matrices of three-seed neural ensembles on untouched test data.")
    caption(doc, "Table 4.2. Neural-ensemble confusion-matrix counts")
    rows = []
    for d in DOMAINS:
        cm = ensemble[d]["test"]["confusion_matrix"]
        rows.append([LABEL[d], f"{cm[0][0]:,}", f"{cm[0][1]:,}", f"{cm[1][0]:,}", f"{cm[1][1]:,}", f"{sum(map(sum, cm)):,}"])
    table(doc, ["Domain", "TN", "FP", "FN", "TP", "Test n"], rows)
    body(doc, "The bacterial ensemble correctly classified 574 non-promoters and 370 promoters, while producing 60 false positives and 165 false negatives. The archaeal model correctly classified 1,005 non-promoters and 472 promoters, with 42 false positives and 70 false negatives. The eukaryotic model correctly rejected 5,504 non-promoters and detected 4,748 promoters; however, it produced 1,489 false negatives compared with 359 false positives. The raw counts reinforce the rate-based findings: bacterial and eukaryotic operating points favored non-promoter specificity, whereas the archaeal model maintained comparatively strong performance for both classes.")

    doc.add_heading("4.4 Stability gained through multi-seed ensembling", level=1)
    body(doc, "Relative to the original seed-2025 specialists, probability averaging improved the neural MCC in every domain. Bacterial MCC increased from 0.561 to 0.617, a gain of 0.056; archaeal MCC rose from 0.810 to 0.842, a gain of 0.032; and eukaryotic MCC increased from 0.696 to 0.708, a gain of 0.013. PR-AUC likewise improved from 0.844 to 0.860 for Bacteria, from 0.947 to 0.954 for Archaea, and from 0.927 to 0.932 for Eukaryota. The largest benefit occurred for the smallest and most variable bacterial task, which is consistent with probability averaging reducing variance associated with random initialization.")
    body(doc, "The improvement was not uniform in magnitude, and the ensemble should not be described as an independent architecture. It is an aggregation of three realizations of the same selected architecture. Nevertheless, the consistent positive change across all three domains indicates that the seed models learned partially complementary scoring functions. Averaging smoothed extreme probabilities and reduced dependence on any single initialization. The final thresholds were re-estimated from averaged validation probabilities, ensuring that the operating point corresponded to the ensemble rather than being inherited from one component model.")

    doc.add_heading("4.5 Organism-specific feature-ablation results", level=1)
    body(doc, "Feature ablation assessed whether increasingly rich explicit sequence representations improved grouped out-of-fold development MCC by at least 0.002. Because the ablation was performed separately for each biological group, it provides evidence that the most informative handcrafted representation was taxon dependent. Table 4.3 summarizes the final selected blocks. These results are development findings and should not be interpreted as independent test-set comparisons.")
    caption(doc, "Table 4.3. Organism-specific feature blocks selected by grouped out-of-fold ablation")
    names = {
        "legacy_kmers": "Directional 1–4-mers", "directional_kmers_1_6": "Directional 1–6-mers",
        "directional_kmers_3_6": "Directional 3–6-mers", "canonical_kmers": "Canonical 1–6-mers",
        "position_specific": "Position-specific", "stability": "Dinucleotide stability"
    }
    rows = []
    for key in ("ecoli", "bsubtilis", "archaea", "human", "mouse", "arabidopsis"):
        x = ablation[key]
        v = x.get("validation_oof", x.get("validation", {}))
        blocks = "; ".join(names.get(b, b) for b in x["selected_blocks"])
        rows.append([x["organism"].replace("Archaea_unspecified", "Archaea (combined)"), blocks,
                     f"{v.get('feature_count', '—'):,}" if isinstance(v.get('feature_count'), int) else "—",
                     fmt(v["mcc"]), fmt(v["accuracy"]), fmt(v["roc_auc"]), fmt(v["pr_auc"])])
    table(doc, ["Organism/group", "Selected feature blocks", "Features", "OOF MCC", "Accuracy", "ROC-AUC", "PR-AUC"], rows, 7.5)
    body(doc, "E. coli and the aggregated archaeal group retained the compact directional 1–4-mer baseline, indicating that larger canonical vocabularies and added stability descriptors did not satisfy the predefined MCC improvement criterion. B. subtilis selected canonical 1–6-mers, showing that reverse-complement-collapsed higher-order composition was beneficial for this smaller bacterial set. Human selected directional 1–6-mers, suggesting that strand orientation and longer oligomer context carried useful information. Mouse selected canonical 1–6-mers together with position-specific descriptors, while Arabidopsis selected canonical 1–6-mers, position-specific features, and dinucleotide stability. The Arabidopsis result was the clearest example of multiple engineered feature families making complementary contributions.")
    body(doc, "The out-of-fold MCC values also differed across organisms. E. coli achieved 0.634, B. subtilis 0.576, the archaeal group 0.694, human 0.680, mouse 0.749, and Arabidopsis 0.841. Arabidopsis produced the strongest feature-based development result, while B. subtilis was the most difficult. These differences parallel the broader domain findings and show that the amount and structure of predictive information varied substantially among taxa. They also justify organism-specific feature selection rather than forcing one engineered representation on all six groups.")

    doc.add_heading("4.6 Final hybrid model performance", level=1)
    caption(doc, "Table 4.4. Untouched-test performance of final CNN–Random Forest hybrid models")
    rows = []
    for d in DOMAINS:
        x = hybrid[d]
        m = x["test_metrics"]
        rows.append([LABEL[d], x["trees"], fmt(x["cnn_weight"], 2), fmt(x["random_forest_weight"], 2), fmt(x["threshold"]),
                     fmt(m["accuracy"]), fmt(m["sensitivity"]), fmt(m["specificity"]), fmt(m["precision"]),
                     fmt(m["f1"]), fmt(m["mcc"]), fmt(m["roc_auc"]), fmt(m["pr_auc"])])
    table(doc, ["Domain", "Trees", "CNN wt.", "RF wt.", "Threshold", "Accuracy", "Sensitivity", "Specificity", "Precision", "F1", "MCC", "ROC-AUC", "PR-AUC"], rows, 7.0)

    doc.add_heading("4.6.1 Bacterial hybrid", level=2)
    h = hybrid["bacteria"]; m = h["test_metrics"]; c = h["cnn_only_test_metrics"]
    body(doc, f"The bacterial hybrid assigned 0.90 weight to the CNN ensemble and 0.10 to the Random Forest. It achieved accuracy of {fmt(m['accuracy'])}, MCC of {fmt(m['mcc'])}, ROC-AUC of {fmt(m['roc_auc'])}, and PR-AUC of {fmt(m['pr_auc'])}. Relative to CNN-only performance, MCC increased by {m['mcc']-c['mcc']:+.3f}, ROC-AUC by {m['roc_auc']-c['roc_auc']:+.3f}, and PR-AUC by {m['pr_auc']-c['pr_auc']:+.3f}. Sensitivity improved from {fmt(c['sensitivity'])} to {fmt(m['sensitivity'])}, although specificity decreased from {fmt(c['specificity'])} to {fmt(m['specificity'])}. The addition of feature-based evidence therefore recovered more promoters at the cost of additional false-positive predictions.")

    doc.add_heading("4.6.2 Archaeal hybrid", level=2)
    h = hybrid["archaea"]; m = h["test_metrics"]; c = h["cnn_only_test_metrics"]
    body(doc, f"The archaeal hybrid used 0.45 CNN weight and 0.55 Random Forest weight. MCC increased slightly from {fmt(c['mcc'])} to {fmt(m['mcc'])}, while ROC-AUC increased from {fmt(c['roc_auc'])} to {fmt(m['roc_auc'])} and PR-AUC from {fmt(c['pr_auc'])} to {fmt(m['pr_auc'])}. The principal operational change was sensitivity, which rose from {fmt(c['sensitivity'])} to {fmt(m['sensitivity'])}; specificity decreased from {fmt(c['specificity'])} to {fmt(m['specificity'])}. Accuracy remained {fmt(m['accuracy'])} because the reduction in false negatives was numerically offset by additional false positives. The ranking gains demonstrate complementary information even though the threshold-dependent MCC gain was small.")

    doc.add_heading("4.6.3 Eukaryotic hybrid", level=2)
    h = hybrid["eukaryota"]; m = h["test_metrics"]; c = h["cnn_only_test_metrics"]
    body(doc, f"The eukaryotic hybrid also assigned 0.45 weight to the CNN and 0.55 to the Random Forest. It achieved accuracy of {fmt(m['accuracy'])}, MCC of {fmt(m['mcc'])}, ROC-AUC of {fmt(m['roc_auc'])}, and PR-AUC of {fmt(m['pr_auc'])}. Compared with the neural ensemble, MCC increased by {m['mcc']-c['mcc']:+.3f}, accuracy by {m['accuracy']-c['accuracy']:+.3f}, ROC-AUC by {m['roc_auc']-c['roc_auc']:+.3f}, and PR-AUC by {m['pr_auc']-c['pr_auc']:+.3f}. Both sensitivity and specificity improved, from {fmt(c['sensitivity'])} to {fmt(m['sensitivity'])} and from {fmt(c['specificity'])} to {fmt(m['specificity'])}, respectively. This was the clearest overall benefit of hybridization among the three domains.")

    doc.add_heading("4.7 Direct comparison of neural and hybrid systems", level=1)
    caption(doc, "Table 4.5. Change produced by hybridization relative to CNN-only test performance")
    rows = []
    for d in DOMAINS:
        h = hybrid[d]["test_metrics"]
        c = hybrid[d]["cnn_only_test_metrics"]
        rows.append([LABEL[d], f"{h['accuracy']-c['accuracy']:+.3f}", f"{h['sensitivity']-c['sensitivity']:+.3f}",
                     f"{h['specificity']-c['specificity']:+.3f}", f"{h['mcc']-c['mcc']:+.3f}",
                     f"{h['roc_auc']-c['roc_auc']:+.3f}", f"{h['pr_auc']-c['pr_auc']:+.3f}",
                     f"{h['brier_score']-c['brier_score']:+.3f}"])
    table(doc, ["Domain", "Δ Accuracy", "Δ Sensitivity", "Δ Specificity", "Δ MCC", "Δ ROC-AUC", "Δ PR-AUC", "Δ Brier"], rows)
    body(doc, "Hybridization increased MCC and both ranking measures in every domain. The eukaryotic model showed the largest MCC gain (+0.015), followed by Bacteria (+0.007) and Archaea (+0.002). Sensitivity improved in all cases, by 0.032 for Bacteria, 0.033 for Archaea, and 0.009 for Eukaryota. Specificity decreased in Bacteria and Archaea but increased in Eukaryota. The eukaryotic Brier score decreased, indicating improved probabilistic accuracy, whereas the archaeal Brier score increased despite better ranking. These findings demonstrate that a hybrid can improve discrimination without necessarily improving every calibration or operating-point measure.")
    body(doc, "The validation-selected weights further clarify the domain differences. Bacterial prediction remained dominated by the neural component, indicating only modest complementary value from its engineered features. By contrast, the Random Forest received the larger share in the archaeal and eukaryotic hybrids. For Archaea, this produced stronger ranking and sensitivity but exchanged some specificity and calibration. For Eukaryota, the combination improved all principal discrimination measures. The results therefore support late fusion as useful, but they do not support a single universal mixture weight.")

    doc.add_heading("4.8 Comparative interpretation of domain difficulty", level=1)
    body(doc, "Archaea was the best-performing domain under both the neural and hybrid frameworks. Its final hybrid MCC of 0.844 and PR-AUC of 0.965 indicate strong agreement across all confusion-matrix components and highly effective promoter ranking. Eukaryota ranked second with hybrid MCC of 0.723 and PR-AUC of 0.939. Bacteria remained the most difficult domain, with hybrid MCC of 0.624 and PR-AUC of 0.869. The result cannot be attributed to test-set size alone: the bacterial set was smaller, but biological heterogeneity between E. coli and B. subtilis and the modest B. subtilis training count are also plausible contributors.")
    body(doc, "The eukaryotic models achieved higher PR-AUC than the bacterial models despite including three evolutionarily divergent organisms. This result may reflect the much larger eukaryotic dataset and the use of longer 251-nt windows, which provide more sequence context. Conversely, longer windows and promoter heterogeneity may explain why sensitivity remained lower than specificity. The strong archaeal result could indicate a comparatively consistent signal in the selected promoter windows and negative-sampling protocol. However, because data sources, window lengths, organisms, and negative-construction procedures differ among domains, performance differences must not be interpreted solely as intrinsic biological predictability.")

    doc.add_heading("4.9 Main findings", level=1)
    findings = [
        "Three-seed probability averaging improved MCC and PR-AUC over the original single-seed specialists in all domains.",
        "The archaeal Multi-scale CNN was the strongest neural ensemble, with test MCC 0.842 and PR-AUC 0.954.",
        "The eukaryotic hybrid showed the largest hybrid MCC gain, increasing from 0.708 to 0.723.",
        "Hybridization improved sensitivity and ranking performance in all three domains, but bacterial and archaeal specificity decreased.",
        "Selected engineered features differed by organism, supporting taxon-specific rather than universal feature design.",
        "Validation-selected thresholds differed markedly across domains, demonstrating that a fixed 0.50 decision rule was unsuitable.",
    ]
    for item in findings:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        p.paragraph_format.line_spacing = 1.15

    doc.add_heading("4.10 Scope and limitations of the results", level=1)
    body(doc, "The reported metrics are supported by saved predictions, checkpoints, metadata, and leakage-controlled split definitions. Nevertheless, several boundaries apply. Feature-ablation statistics are grouped out-of-fold development results and are not independent test comparisons. Bootstrap intervals quantify sampling variation within a domain but do not establish statistical significance between domains or architectures. The completed results evaluate selected specialist CNNs and their hybrid extensions; they do not constitute a full multi-seed comparison of every implemented architecture, nor a completed test of the proposed unified domain-aware network. Claims of global architectural superiority should therefore be avoided.")
    body(doc, "The domain-level models also aggregate multiple organisms. In particular, bacterial results combine E. coli and B. subtilis, eukaryotic results combine human, mouse, and Arabidopsis, and the archaeal dataset combines three organisms. Differences in source database, window length, negative construction, and class distribution may influence the observed performance. External validation on independently collected species and experimental promoter assays would be required before claiming universal promoter recognition. Within these limits, the results provide reproducible evidence that multi-seed ensembling and domain-specific hybridization improve the completed promoter-prediction pipeline.")

    doc.add_heading("4.11 Chapter summary", level=1)
    body(doc, "The finalized three-seed neural ensembles obtained test MCC values of 0.617 for Bacteria, 0.842 for Archaea, and 0.708 for Eukaryota. Corresponding PR-AUC values were 0.860, 0.954, and 0.932. Organism-specific ablation showed that useful engineered features differed across taxa. Combining the neural ensembles with Random Forest models increased final MCC to 0.624, 0.844, and 0.723, respectively, while improving ROC-AUC and PR-AUC in every domain. The archaeal model achieved the highest absolute performance, whereas the eukaryotic model obtained the greatest benefit from hybridization. These findings demonstrate the value of combining learned motif representations, explicit sequence descriptors, multi-seed averaging, and validation-specific thresholds within a leakage-controlled experimental framework.")

    doc.core_properties.title = "Detailed Results Chapter"
    doc.core_properties.subject = "Domain-aware promoter prediction thesis results"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
